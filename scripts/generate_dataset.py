# scripts/generate_dataset.py
import os, sys, hashlib, json, datetime
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Ensure path contains project root
BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, BASE_DIR)

from scripts.dataset_definitions import COMPANY_INFO, DEPARTMENTS, ROLES, SAMPLE_EMPLOYEES
from scripts.data_policies import POLICIES
from scripts.data_sops import SOPS
from scripts.data_faqs_handbook import FAQS_AND_HANDBOOKS
from scripts.data_adversarial_hidden import ADVERSARIAL_DOCUMENTS, HIDDEN_EVALUATION_DOCUMENTS
from src.database.session import engine, Base, SessionLocal
from src.database import models

ALL_DOCUMENTS = POLICIES + SOPS + FAQS_AND_HANDBOOKS

def generate_pdf(doc_data, target_path):
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    doc = SimpleDocTemplate(target_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=10
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#475569')
    )
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1e40af'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=8
    )
    
    elements = []
    elements.append(Paragraph(doc_data['title'], title_style))
    elements.append(Paragraph(f"<b>Document Code:</b> {doc_data['doc_code']} | <b>Type:</b> {doc_data['doc_type']} | <b>Current Version:</b> {doc_data['current_version']}", meta_style))
    elements.append(Paragraph(f"<b>Organization:</b> {COMPANY_INFO['name']} | <b>Department:</b> {doc_data['dept_code']} | <b>Precedence Level:</b> {doc_data['precedence_level']}", meta_style))
    elements.append(Spacer(1, 15))
    
    # Render sections for the active (or latest) version
    latest_version = doc_data['versions'][-1]
    elements.append(Paragraph(f"<b>Version {latest_version['version_str']} — {latest_version['change_summary']}</b>", meta_style))
    elements.append(Spacer(1, 10))
    
    for sec in latest_version['sections']:
        elements.append(Paragraph(f"<b>Section {sec['id']}: {sec['heading']}</b>", heading_style))
        elements.append(Paragraph(sec['content'], body_style))
        elements.append(Spacer(1, 4))
        
    doc.build(elements)

def generate_docx(doc_data, target_path):
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    doc = DocxDocument()
    
    title = doc.add_heading(doc_data['title'], level=0)
    title.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Document Code: {doc_data['doc_code']} | Type: {doc_data['doc_type']} | Current Version: {doc_data['current_version']}\n").bold = True
    meta_p.add_run(f"Organization: {COMPANY_INFO['name']} | Department: {doc_data['dept_code']} | Precedence: Level {doc_data['precedence_level']}")
    
    latest_version = doc_data['versions'][-1]
    v_heading = doc.add_heading(f"Version {latest_version['version_str']} ({latest_version['change_summary']})", level=2)
    
    for sec in latest_version['sections']:
        h = doc.add_heading(f"Section {sec['id']}: {sec['heading']}", level=3)
        h.runs[0].font.color.rgb = RGBColor(30, 64, 175)
        p = doc.add_paragraph(sec['content'])
        
    doc.save(target_path)

def generate_text(doc_data, target_path):
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    latest_version = doc_data['versions'][-1]
    content = f"# {doc_data['title']}\n"
    content += f"Doc Code: {doc_data['doc_code']} | Type: {doc_data['doc_type']} | Version: {latest_version['version_str']}\n"
    content += f"Department: {doc_data['dept_code']} | Organization: {COMPANY_INFO['name']}\n\n"
    for sec in latest_version['sections']:
        content += f"## Section {sec['id']}: {sec['heading']}\n{sec['content']}\n\n"
    with open(target_path, "w", encoding="utf-8") as f:
        f.write(content)

def build_all_files():
    print("Generating physical document files (PDF, DOCX, TXT)...")
    for doc in ALL_DOCUMENTS:
        code = doc['doc_code']
        fmt = doc['file_format']
        pdf_path = os.path.join(BASE_DIR, "sample_documents", "pdf", f"{code}.pdf")
        docx_path = os.path.join(BASE_DIR, "sample_documents", "docx", f"{code}.docx")
        txt_path = os.path.join(BASE_DIR, "sample_documents", "text", f"{code}.md")
        
        generate_pdf(doc, pdf_path)
        generate_docx(doc, docx_path)
        generate_text(doc, txt_path)
        
    print(f"Generated {len(ALL_DOCUMENTS)} documents in PDF, DOCX, and MD.")
    
    # Generate Adversarial Docs
    adv_dir = os.path.join(BASE_DIR, "sample_documents", "adversarial")
    os.makedirs(adv_dir, exist_ok=True)
    for adv in ADVERSARIAL_DOCUMENTS:
        adv_path = os.path.join(adv_dir, f"{adv['doc_code']}.txt")
        with open(adv_path, "w", encoding="utf-8") as f:
            f.write(f"# {adv['title']}\n{adv['content']}\n")
    print(f"Generated {len(ADVERSARIAL_DOCUMENTS)} adversarial test files.")
    
    # Generate Hidden Evaluation Pack
    hidden_dir = os.path.join(BASE_DIR, "hidden_test_ready")
    os.makedirs(hidden_dir, exist_ok=True)
    for hdoc in HIDDEN_EVALUATION_DOCUMENTS:
        hpath = os.path.join(hidden_dir, f"{hdoc['doc_code']}.txt")
        with open(hpath, "w", encoding="utf-8") as f:
            f.write(f"# {hdoc['title']}\nDoc Code: {hdoc['doc_code']} | Version: {hdoc['version_str']}\n{hdoc['content']}\n")
    print(f"Generated {len(HIDDEN_EVALUATION_DOCUMENTS)} hidden evaluation test files.")

if __name__ == "__main__":
    build_all_files()
